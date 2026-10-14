"""Extract plain text from an existing CV (.docx, .pdf, .md, .txt) so it can be restructured into master_cv.md.

Usage: python import_cv.py <file> [--out text.md]
Prints to stdout when --out is omitted.
"""
import argparse
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def from_docx(path):
    try:
        from docx import Document
    except ImportError:
        raise SystemExit("python-docx is not installed. Run: pip install python-docx")
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            lines.append("")
            continue
        style = (p.style.name or "").lower()
        if "heading 1" in style or "title" in style:
            lines.append("# " + text)
        elif "heading" in style:
            lines.append("## " + text)
        elif "list" in style:
            lines.append("* " + text)
        else:
            lines.append(text)
    for table in doc.tables:
        lines.append("")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def from_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise SystemExit("pypdf is not installed. Run: pip install pypdf")
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--out")
    args = ap.parse_args()
    path = os.path.abspath(args.file)
    if not os.path.exists(path):
        print("ERROR: file not found: " + path)
        return 1
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        text = from_docx(path)
    elif ext == ".pdf":
        text = from_pdf(path)
    elif ext in (".md", ".txt", ".markdown"):
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
    else:
        print("ERROR: unsupported extension " + ext + " (use .docx, .pdf, .md or .txt)")
        return 1
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(text)
        print("WROTE: " + os.path.abspath(args.out) + " (" + str(len(text)) + " chars)")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
