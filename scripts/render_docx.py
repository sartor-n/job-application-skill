"""Parsed CV (or a simple markdown document) -> .docx via python-docx, using the same theme as the HTML."""
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from cv_parser import parse_inline, HEADER_RE, BULLET_RE, SEPARATOR_RE, split_frontmatter  # noqa: E402
from render_common import contact_parts, display_name, display_title, PAGE_SIZES_MM  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor, Emu
    HAVE_DOCX = True
except ImportError:  # pragma: no cover
    HAVE_DOCX = False


def _rgb(hex_colour):
    return RGBColor.from_string(hex_colour.lstrip("#").upper()[:6])


def _first_font(family):
    return family.split(",")[0].strip().strip("'\"") or "Calibri"


def _set_font(run, name, size_pt=None, bold=None, italic=None, colour=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if colour:
        run.font.color.rgb = _rgb(colour)


def _add_hyperlink(paragraph, url, text, font, size_pt, colour):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), colour.lstrip("#").upper())
    rpr.append(color)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _bottom_border(paragraph, colour, size=6):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), colour.lstrip("#").upper())
    pbdr.append(bottom)
    ppr.append(pbdr)


class DocxWriter:
    def __init__(self, theme):
        if not HAVE_DOCX:
            raise RuntimeError("python-docx is not installed. Run: pip install python-docx")
        self.theme = theme
        self.font = _first_font(theme["font_family"])
        self.body_pt = float(theme["body_pt"])
        self.doc = Document()
        sec = self.doc.sections[0]
        w, h = PAGE_SIZES_MM[theme["page_size"]]
        sec.page_width, sec.page_height = Mm(w), Mm(h)
        t, r, b, l = theme["margins_mm"]
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = Mm(t), Mm(r), Mm(b), Mm(l)
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        normal = self.doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(self.body_pt)
        normal.font.color.rgb = _rgb(theme["text_colour"])
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)
        normal.paragraph_format.line_spacing = float(theme["line_spacing"])
        normal.paragraph_format.space_after = Pt(float(theme["para_after_pt"]))
        normal.paragraph_format.space_before = Pt(0)
        for style_name in ("List Bullet",):
            st = self.doc.styles[style_name]
            st.font.name = self.font
            st.font.size = Pt(self.body_pt)
            st.paragraph_format.space_after = Pt(2)
            st.paragraph_format.line_spacing = float(theme["line_spacing"])

    def runs(self, paragraph, runs, size_pt=None, base_bold=False, colour=None):
        size = size_pt or self.body_pt
        for r in runs:
            if r["href"]:
                _add_hyperlink(paragraph, r["href"], r["text"], self.font, size, self.theme["accent"])
                continue
            run = paragraph.add_run(r["text"])
            _set_font(run, self.font, size, bold=(r["bold"] or base_bold), italic=r["italic"], colour=colour)

    def inline(self, paragraph, text, **kw):
        self.runs(paragraph, parse_inline(text), **kw)

    def header(self, profile, cv, market):
        name = display_name(profile, cv)
        title = display_title(profile, cv)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        _set_font(p.add_run(name), self.font, 17, bold=True, colour="#0f172a")
        if title:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            self.inline(p, title, size_pt=11, base_bold=True, colour=self.theme["heading_colour"])
        parts = contact_parts(profile, cv, market)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        for i, (text, href) in enumerate(parts):
            if i:
                _set_font(p.add_run("  |  "), self.font, 9, colour="#94a3b8")
            if href:
                _add_hyperlink(p, href, text, self.font, 9, self.theme["accent"])
            else:
                _set_font(p.add_run(text), self.font, 9, colour="#475569")
        _bottom_border(p, self.theme["accent"], size=18)

    def section_heading(self, title):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        _set_font(p.add_run(title.upper()), self.font, 10.5, bold=True, colour=self.theme["heading_colour"])
        _bottom_border(p, "#cbd5e1", size=4)

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        self.inline(p, text)

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            self.inline(p, item)

    def block(self, b):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Emu(self.text_width), WD_TAB_ALIGNMENT.RIGHT)
        self.inline(p, b["role"], size_pt=10.5, base_bold=True, colour="#0f172a")
        if b["org"]:
            _set_font(p.add_run(" | "), self.font, 10.5, bold=True, colour="#334155")
            self.inline(p, b["org"], size_pt=10.5, base_bold=True, colour="#334155")
        if b["date"]:
            _set_font(p.add_run("\t" + b["date"]), self.font, 9, italic=True, colour="#64748b")
        for para in b["paragraphs"]:
            self.paragraph(para)
        self.bullets(b["items"])

    def section(self, sec):
        self.section_heading(sec["title"])
        for para in sec["paragraphs"]:
            self.paragraph(para)
        self.bullets(sec["items"])
        for b in sec["blocks"]:
            self.block(b)

    def markdown_doc(self, md_text):
        _, body = split_frontmatter(md_text)
        para = []

        def flush():
            nonlocal para
            if para:
                self.paragraph(" ".join(para))
                para = []

        for raw in body.split("\n"):
            line = raw.strip()
            if not line or SEPARATOR_RE.match(line):
                flush()
                continue
            if line.startswith("<!--") or line.endswith("-->"):
                continue
            hm = HEADER_RE.match(line)
            if hm:
                flush()
                lvl = len(hm.group(1))
                p = self.doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8 if lvl <= 2 else 5)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                size = {1: 14, 2: 11.5}.get(lvl, 10.5)
                colour = self.theme["heading_colour"] if lvl <= 2 else "#0f172a"
                self.inline(p, hm.group(2), size_pt=size, base_bold=True, colour=colour)
                continue
            bm = BULLET_RE.match(line)
            if bm:
                flush()
                self.bullets([bm.group(1)])
                continue
            para.append(line)
        flush()

    def save(self, path):
        self.doc.save(path)


def render_cv_docx(cv, profile, theme, market, out_path):
    w = DocxWriter(theme)
    w.header(profile, cv, market)
    for sec in cv["sections"]:
        w.section(sec)
    w.save(out_path)
    return out_path


def render_doc_docx(md_text, profile, cv, theme, market, out_path):
    w = DocxWriter(theme)
    w.header(profile, cv, market)
    w.markdown_doc(md_text)
    w.save(out_path)
    return out_path
